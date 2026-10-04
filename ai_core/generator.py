"""Formatting utilities: sanitize, DOCX, PDF, HTML preview."""
import html
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from config import FOOTER_TEXT, LOGO_PATH

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u00a0": " ",
    "\u2022": "-",
}


def sanitize_text(text: str) -> str:
    """Replace typographic characters and strip markdown noise."""
    for k, v in _REPLACEMENTS.items():
        text = text.replace(k, v)
    text = re.sub(r"```\w*", "", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    return text.strip()


def split_terms(terms: str) -> list:
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


def is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    return bool(
        re.match(r"^\d+(\.\d+)*[.)]\s+\S", s) and s.endswith(":")
        or re.match(r"^\d+\.\s+[A-Z][^.]{0,60}:?$", s)
        or s.endswith(":") and len(s) < 60
        or s.isupper() and len(s) > 3
    )


def _lines(text: str):
    return [ln.rstrip() for ln in text.splitlines()]


# ---------------------------------------------------------------- DOCX
def _add_footer(section, text):
    p = section.footer.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.name = "Times New Roman"


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    text = sanitize_text(text)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(2.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(doc_type.strip().title())
    r.bold = True
    r.font.size = Pt(16)

    for line in _lines(text):
        if not line.strip():
            continue
        para = doc.add_paragraph()
        if is_heading(line):
            para.add_run(line.strip()).bold = True
        elif line.lstrip().startswith(("- ", "* ")):
            para.style = "List Bullet"
            para.add_run(line.lstrip()[2:])
        else:
            para.add_run(line.strip())
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    items = split_terms(terms)
    if items:
        doc.add_paragraph().add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "No.", "Term"
        for c in hdr:
            _shade(c, "D9D9D9")
            for r in c.paragraphs[0].runs:
                r.bold = True
        for i, t in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text, row[1].text = str(i), sanitize_text(t)
        for row in table.rows:
            row.cells[0].width = Inches(0.6)
            row.cells[1].width = Inches(5.6)

    _add_footer(doc.sections[0], FOOTER_TEXT)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------- PDF
def _latin(s: str) -> str:
    return sanitize_text(s).encode("latin-1", "replace").decode("latin-1")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    from fpdf import FPDF  # imported lazily so the rest works without fpdf

    class LegalPDF(FPDF):
        def header(self):
            if LOGO_PATH.exists():
                self.image(str(LOGO_PATH), x=(self.w - 45) / 2, y=8, w=45)
            self.set_y(30)
            self.set_font("Helvetica", "B", 13)
            self.cell(0, 8, _latin(doc_type.title()), align="C")
            self.ln(12)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 8, _latin(FOOTER_TEXT), align="C")

    pdf = LegalPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.set_margins(20, 20, 20)
    pdf.add_page()

    for line in _lines(_latin(text)):
        if not line.strip():
            pdf.ln(3)
            continue
        if is_heading(line):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 6, line.strip(), new_x="LMARGIN", new_y="NEXT") \
                if _is_fpdf2() else pdf.multi_cell(0, 6, line.strip())
            pdf.ln(1)
        else:
            body = line.strip()
            if body.startswith(("- ", "* ")):
                body = "  -  " + body[2:]
            pdf.set_font("Helvetica", "", 10.5)
            pdf.multi_cell(0, 5.5, body, align="J") if not _is_fpdf2() else \
                pdf.multi_cell(0, 5.5, body, align="J", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)

    items = split_terms(terms)
    if items:
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 11)
        _mc(pdf, 6, "Summary of Key Terms")
        pdf.set_font("Helvetica", "", 10.5)
        for t in items:
            _mc(pdf, 5.5, "  -  " + _latin(t))

    out = pdf.output(dest="S") if not _is_fpdf2() else pdf.output()
    return out.encode("latin-1") if isinstance(out, str) else bytes(out)


def _is_fpdf2() -> bool:
    import fpdf
    return hasattr(fpdf, "FPDF") and getattr(fpdf, "__version__", "1")[0] >= "2"


def _mc(pdf, h, txt):
    if _is_fpdf2():
        pdf.multi_cell(0, h, txt, new_x="LMARGIN", new_y="NEXT")
    else:
        pdf.multi_cell(0, h, txt)


# ---------------------------------------------------------------- HTML
def format_html_preview(text: str) -> str:
    out = []
    for line in _lines(sanitize_text(text)):
        if not line.strip():
            out.append("<br>")
            continue
        esc = html.escape(line.strip())
        if is_heading(line):
            out.append(f"<h4 style='margin:14px 0 4px;color:#fff'>{esc}</h4>")
        elif line.lstrip().startswith(("- ", "* ")):
            out.append(f"<li>{html.escape(line.lstrip()[2:])}</li>")
        else:
            out.append(f"<p style='margin:4px 0'>{esc}</p>")
    return "\n".join(out)
